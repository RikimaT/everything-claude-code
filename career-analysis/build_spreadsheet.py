#!/usr/bin/env python3
"""
build_spreadsheet.py

Generate the career analysis spreadsheet (.xlsx) for Bridge Tutoring School.
Creates 4 new sheets to be imported into the existing Google Spreadsheet:
  1. Middle school correction master
  2. Conversion master
  3. Test score input & school prediction
  4. OCR read data storage

All prediction logic uses spreadsheet formulas only (no macros/GAS).
"""

import os
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# --- Constants ---
OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "ブリッジ進路分析システム_新規シート.xlsx")

# Styling
HEADER_FONT = Font(name="Meiryo", bold=True, size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT_WHITE = Font(name="Meiryo", bold=True, size=11, color="FFFFFF")
CELL_FONT = Font(name="Meiryo", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
NOTE_FONT = Font(name="Meiryo", size=9, italic=True, color="808080")

# Judgment colors
JUDGMENT_COLORS = {
    "A": PatternFill(start_color="CFE2F3", end_color="CFE2F3", fill_type="solid"),  # Blue
    "B": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),  # Green
    "C": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),  # Yellow
    "D": PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid"),  # Orange
    "E": PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid"),  # Red
}

# Middle schools
MIDDLE_SCHOOLS = [
    ("稲枝中", 1.00, "基準校"),
    ("彦根東中", None, "田中氏に確認後入力"),
    ("彦根中央中", None, "田中氏に確認後入力"),
    ("南中", None, "田中氏に確認後入力"),
    ("彦根西中", None, "田中氏に確認後入力"),
    ("鳥居本中", None, "田中氏に確認後入力"),
    ("多賀中", None, "田中氏に確認後入力"),
    ("河瀬中", None, "田中氏に確認後入力"),
]

PREDICTION_ROWS = 30  # Number of data rows in prediction sheet
OCR_ROWS = 100  # Number of data rows in OCR sheet
CONVERSION_ROWS = 20  # Number of data rows in conversion master


def apply_header_style(cell):
    """Apply header styling to a cell."""
    cell.font = HEADER_FONT_WHITE
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


def apply_cell_style(cell):
    """Apply standard cell styling."""
    cell.font = CELL_FONT
    cell.border = THIN_BORDER
    cell.alignment = Alignment(vertical="center")


def set_column_width(ws, col_letter, width):
    """Set column width."""
    ws.column_dimensions[col_letter].width = width


# =============================================================================
# Sheet 1: Middle School Correction Master
# =============================================================================
def create_school_correction_master(wb):
    """Create the middle school correction master sheet."""
    ws = wb.create_sheet(title="中学校補正マスタ")

    # Headers
    headers = ["中学校名", "補正係数", "根拠・メモ"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)

    # Data
    for row_idx, (school, coeff, memo) in enumerate(MIDDLE_SCHOOLS, 2):
        cell_a = ws.cell(row=row_idx, column=1, value=school)
        apply_cell_style(cell_a)

        cell_b = ws.cell(row=row_idx, column=2)
        if coeff is not None:
            cell_b.value = coeff
            cell_b.number_format = "0.00"
        apply_cell_style(cell_b)

        cell_c = ws.cell(row=row_idx, column=3, value=memo)
        apply_cell_style(cell_c)

    # Column widths
    set_column_width(ws, "A", 16)
    set_column_width(ws, "B", 14)
    set_column_width(ws, "C", 30)

    # Freeze header row
    ws.freeze_panes = "A2"

    return ws


# =============================================================================
# Sheet 2: Conversion Master
# =============================================================================
def create_conversion_master(wb):
    """Create the conversion master sheet (corrected score -> estimated deviation)."""
    ws = wb.create_sheet(title="換算マスタ")

    # Headers
    headers = ["補正後5科合計（下限）", "推定Vもし偏差値"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        apply_header_style(cell)

    # Empty data rows with styling
    for row_idx in range(2, CONVERSION_ROWS + 2):
        for col in range(1, 3):
            cell = ws.cell(row=row_idx, column=col)
            apply_cell_style(cell)
            cell.number_format = "0.0" if col == 2 else "0"

    # Note about sort order
    note_row = CONVERSION_ROWS + 3
    note_cell = ws.cell(
        row=note_row, column=1,
        value="※ A列（補正後5科合計）は必ず昇順で入力してください（VLOOKUPの近似一致で正しく動作するため）"
    )
    note_cell.font = NOTE_FONT

    # Column widths
    set_column_width(ws, "A", 24)
    set_column_width(ws, "B", 20)

    # Freeze header row
    ws.freeze_panes = "A2"

    return ws


# =============================================================================
# Sheet 3: Test Score Input & School Prediction (Main Sheet)
# =============================================================================
def create_prediction_sheet(wb):
    """Create the main prediction sheet with formulas, validation, and formatting."""
    ws = wb.create_sheet(title="テスト点数入力＆進学先予測")

    # --- Headers ---
    headers = [
        ("A", "生徒名"),
        ("B", "学年"),
        ("C", "中学校名"),
        ("D", "国語"),
        ("E", "数学"),
        ("F", "英語"),
        ("G", "理科"),
        ("H", "社会"),
        ("I", "5科合計"),
        ("J", "Vもし偏差値"),
        ("K", "補正係数"),
        ("L", "補正後5科合計"),
        ("M", "推定Vもし偏差値"),
        ("N", "志望校"),
        ("O", "志望校偏差値"),
        ("P", "判定"),
        ("Q", "予測第1候補"),
        ("R", "予測第2候補"),
        ("S", "予測第3候補"),
    ]

    for col_letter, header in headers:
        col_idx = ord(col_letter) - ord("A") + 1
        cell = ws.cell(row=1, column=col_idx, value=header)
        apply_header_style(cell)

    # --- Column widths ---
    col_widths = {
        "A": 14, "B": 8, "C": 14, "D": 8, "E": 8, "F": 8, "G": 8, "H": 8,
        "I": 10, "J": 14, "K": 10, "L": 14, "M": 16, "N": 16, "O": 14,
        "P": 18, "Q": 16, "R": 16, "S": 16,
    }
    for col_letter, width in col_widths.items():
        set_column_width(ws, col_letter, width)

    # --- Formulas and styling for each data row ---
    for r in range(2, PREDICTION_ROWS + 2):
        # A: Student name (manual input)
        apply_cell_style(ws.cell(row=r, column=1))

        # B: Grade (dropdown set below)
        apply_cell_style(ws.cell(row=r, column=2))

        # C: Middle school name (dropdown set below)
        apply_cell_style(ws.cell(row=r, column=3))

        # D-H: Subject scores (manual input, validation set below)
        for col in range(4, 9):
            apply_cell_style(ws.cell(row=r, column=col))

        # I: 5-subject total (auto-calculated)
        cell_i = ws.cell(row=r, column=9)
        cell_i.value = f'=IF(D{r}<>"",SUM(D{r}:H{r}),"")'
        apply_cell_style(cell_i)

        # J: V-moshi deviation (manual input, optional)
        apply_cell_style(ws.cell(row=r, column=10))

        # K: Correction coefficient (VLOOKUP from correction master)
        cell_k = ws.cell(row=r, column=11)
        cell_k.value = (
            f"=IFERROR(VLOOKUP(C{r},'中学校補正マスタ'!A:B,2,FALSE),\"\")"
        )
        cell_k.number_format = "0.00"
        apply_cell_style(cell_k)

        # L: Corrected 5-subject total
        cell_l = ws.cell(row=r, column=12)
        cell_l.value = (
            f'=IF(AND(I{r}<>"",K{r}<>""),I{r}*K{r},"")'
        )
        cell_l.number_format = "0.0"
        apply_cell_style(cell_l)

        # M: Estimated V-moshi deviation
        cell_m = ws.cell(row=r, column=13)
        cell_m.value = (
            f'=IF(J{r}<>"",J{r},'
            f'IF(L{r}<>"",IFERROR(VLOOKUP(L{r},\'換算マスタ\'!A:B,2,TRUE),""),""))'
        )
        cell_m.number_format = "0.0"
        apply_cell_style(cell_m)

        # N: Target school (dropdown set below)
        apply_cell_style(ws.cell(row=r, column=14))

        # O: Target school deviation (VLOOKUP from high school summary)
        cell_o = ws.cell(row=r, column=15)
        cell_o.value = (
            f"=IFERROR(VLOOKUP(N{r},'高校別サマリー'!A:C,3,FALSE),\"\")"
        )
        cell_o.number_format = "0.0"
        apply_cell_style(cell_o)

        # P: Judgment
        cell_p = ws.cell(row=r, column=16)
        cell_p.value = (
            f'=IF(OR(M{r}="",O{r}=""),"",'
            f'IFS('
            f'M{r}>=O{r}+3,"A判定（安全圏）",'
            f'M{r}>=O{r}+1,"B判定（合格圏）",'
            f'M{r}>=O{r}-1,"C判定（ボーダー）",'
            f'M{r}>=O{r}-3,"D判定（努力圏）",'
            f'TRUE,"E判定（要再考）"))'
        )
        apply_cell_style(cell_p)

        # Q: Prediction 1st candidate
        cell_q = ws.cell(row=r, column=17)
        cell_q.value = (
            f"=IFERROR(INDEX(SORT('高校別サマリー'!A$2:C$19,"
            f"ABS('高校別サマリー'!C$2:C$19-M{r}),TRUE),1,1),\"\")"
        )
        apply_cell_style(cell_q)

        # R: Prediction 2nd candidate
        cell_r = ws.cell(row=r, column=18)
        cell_r.value = (
            f"=IFERROR(INDEX(SORT('高校別サマリー'!A$2:C$19,"
            f"ABS('高校別サマリー'!C$2:C$19-M{r}),TRUE),2,1),\"\")"
        )
        apply_cell_style(cell_r)

        # S: Prediction 3rd candidate
        cell_s = ws.cell(row=r, column=19)
        cell_s.value = (
            f"=IFERROR(INDEX(SORT('高校別サマリー'!A$2:C$19,"
            f"ABS('高校別サマリー'!C$2:C$19-M{r}),TRUE),3,1),\"\")"
        )
        apply_cell_style(cell_s)

    # --- Data Validation ---

    # Grade dropdown (B column)
    dv_grade = DataValidation(
        type="list",
        formula1='"中1,中2,中3"',
        allow_blank=True,
    )
    dv_grade.error = "中1/中2/中3 から選択してください"
    dv_grade.errorTitle = "入力エラー"
    dv_grade.prompt = "学年を選択"
    dv_grade.promptTitle = "学年"
    ws.add_data_validation(dv_grade)
    dv_grade.add(f"B2:B{PREDICTION_ROWS + 1}")

    # Middle school dropdown (C column) - references correction master
    dv_school = DataValidation(
        type="list",
        formula1="'中学校補正マスタ'!$A$2:$A$9",
        allow_blank=True,
    )
    dv_school.error = "中学校補正マスタに登録された中学校名から選択してください"
    dv_school.errorTitle = "入力エラー"
    dv_school.prompt = "中学校名を選択"
    dv_school.promptTitle = "中学校"
    ws.add_data_validation(dv_school)
    dv_school.add(f"C2:C{PREDICTION_ROWS + 1}")

    # Subject score validation (D-H columns, 0-100)
    dv_score = DataValidation(
        type="whole",
        operator="between",
        formula1=0,
        formula2=100,
        allow_blank=True,
    )
    dv_score.error = "0〜100の整数を入力してください"
    dv_score.errorTitle = "入力エラー"
    dv_score.prompt = "0〜100の点数を入力"
    dv_score.promptTitle = "点数"
    ws.add_data_validation(dv_score)
    dv_score.add(f"D2:H{PREDICTION_ROWS + 1}")

    # V-moshi deviation validation (J column, 20-80 reasonable range)
    dv_deviation = DataValidation(
        type="decimal",
        operator="between",
        formula1=20,
        formula2=80,
        allow_blank=True,
    )
    dv_deviation.error = "20〜80の数値を入力してください"
    dv_deviation.errorTitle = "入力エラー"
    dv_deviation.prompt = "Vもし偏差値を入力（空欄可）"
    dv_deviation.promptTitle = "偏差値"
    ws.add_data_validation(dv_deviation)
    dv_deviation.add(f"J2:J{PREDICTION_ROWS + 1}")

    # Target school dropdown (N column) - references high school summary
    dv_target = DataValidation(
        type="list",
        formula1="'高校別サマリー'!$A$2:$A$19",
        allow_blank=True,
    )
    dv_target.error = "高校別サマリーに登録された高校名から選択してください"
    dv_target.errorTitle = "入力エラー"
    dv_target.prompt = "志望校を選択"
    dv_target.promptTitle = "志望校"
    ws.add_data_validation(dv_target)
    dv_target.add(f"N2:N{PREDICTION_ROWS + 1}")

    # --- Conditional Formatting for Judgment column (P) ---
    p_range = f"P2:P{PREDICTION_ROWS + 1}"

    # Use FormulaRule with SEARCH to detect judgment text in cells.
    # The formula references P2 (top of range); Excel/Sheets adjusts for each row.
    judgments = [
        ("A判定", JUDGMENT_COLORS["A"]),
        ("B判定", JUDGMENT_COLORS["B"]),
        ("C判定", JUDGMENT_COLORS["C"]),
        ("D判定", JUDGMENT_COLORS["D"]),
        ("E判定", JUDGMENT_COLORS["E"]),
    ]
    for text, fill in judgments:
        ws.conditional_formatting.add(
            p_range,
            FormulaRule(
                formula=[f'NOT(ISERROR(SEARCH("{text}",P2)))'],
                fill=fill,
            ),
        )

    # Freeze header row
    ws.freeze_panes = "A2"

    return ws


# =============================================================================
# Sheet 4: OCR Read Data
# =============================================================================
def create_ocr_data_sheet(wb):
    """Create the OCR read data storage sheet."""
    ws = wb.create_sheet(title="OCR読取データ")

    # Headers
    headers = [
        ("A", "生徒名"),
        ("B", "学年"),
        ("C", "中学校名"),
        ("D", "データ種別"),
        ("E", "国語"),
        ("F", "数学"),
        ("G", "英語"),
        ("H", "理科"),
        ("I", "社会"),
        ("J", "5科合計"),
        ("K", "Vもし偏差値"),
        ("L", "テスト回/年度"),
        ("M", "読取元ファイル名"),
    ]

    for col_letter, header in headers:
        col_idx = ord(col_letter) - ord("A") + 1
        cell = ws.cell(row=1, column=col_idx, value=header)
        apply_header_style(cell)

    # Column widths
    col_widths = {
        "A": 14, "B": 8, "C": 14, "D": 14, "E": 8, "F": 8, "G": 8,
        "H": 8, "I": 8, "J": 10, "K": 14, "L": 16, "M": 30,
    }
    for col_letter, width in col_widths.items():
        set_column_width(ws, col_letter, width)

    # Data rows with auto-sum formula for column J
    for r in range(2, OCR_ROWS + 2):
        for col in range(1, 14):
            cell = ws.cell(row=r, column=col)
            apply_cell_style(cell)

        # J: 5-subject total (auto-calculated)
        ws.cell(row=r, column=10).value = f'=IF(E{r}<>"",SUM(E{r}:I{r}),"")'

    # Data validation for data type (D column)
    dv_type = DataValidation(
        type="list",
        formula1='"定期テスト,Vもし,進研Sテスト,実力テスト,その他"',
        allow_blank=True,
    )
    dv_type.error = "リストからデータ種別を選択してください"
    dv_type.errorTitle = "入力エラー"
    ws.add_data_validation(dv_type)
    dv_type.add(f"D2:D{OCR_ROWS + 1}")

    # Middle school dropdown (C column)
    dv_school = DataValidation(
        type="list",
        formula1="'中学校補正マスタ'!$A$2:$A$9",
        allow_blank=True,
    )
    ws.add_data_validation(dv_school)
    dv_school.add(f"C2:C{OCR_ROWS + 1}")

    # Freeze header row
    ws.freeze_panes = "A2"

    return ws


# =============================================================================
# Main
# =============================================================================
def main():
    wb = Workbook()

    # Remove the default sheet created by openpyxl
    wb.remove(wb.active)

    # Create all 4 sheets
    create_school_correction_master(wb)
    create_conversion_master(wb)
    create_prediction_sheet(wb)
    create_ocr_data_sheet(wb)

    # Save
    wb.save(OUTPUT_FILE)
    print(f"Generated: {OUTPUT_FILE}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
